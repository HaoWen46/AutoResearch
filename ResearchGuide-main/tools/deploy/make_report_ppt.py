"""Build the 10 Oct 12 team report deck. No secrets."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT = Path(r"C:\Users\Lenovo\Desktop\Self-Evolved-Agent\启研-第3次小组工作报告-2026-10-09.pptx")
INK = RGBColor(0x1F, 0x2A, 0x24)
TEAL = RGBColor(0x1F, 0x6F, 0x64)
MUTED = RGBColor(0x5A, 0x64, 0x5C)
PAPER = RGBColor(0xF5, 0xF2, 0xEA)
LINE = RGBColor(0xD5, 0xCF, 0xC0)


def _set_run(run, text, size=18, bold=False, color=INK):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "微软雅黑"


def add_bg(slide):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    box.line.fill.background()
    box.fill.solid()
    box.fill.fore_color.rgb = PAPER


def kicker(slide, text):
    t = slide.shapes.add_textbox(Inches(0.7), Inches(0.32), Inches(12), Inches(0.35))
    p = t.text_frame.paragraphs[0]
    r = p.add_run()
    _set_run(r, text, 12, True, TEAL)


def title(slide, text):
    t = slide.shapes.add_textbox(Inches(0.7), Inches(0.62), Inches(12), Inches(0.7))
    tf = t.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    _set_run(r, text, 28, True, INK)


def bullets(slide, items, top=1.5, size=18):
    t = slide.shapes.add_textbox(Inches(0.75), Inches(top), Inches(11.8), Inches(5.5))
    tf = t.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(10)
        r = p.add_run()
        _set_run(r, item, size, False, INK)


def footer(slide, n, total=10):
    t = slide.shapes.add_textbox(Inches(0.7), Inches(7.05), Inches(12), Inches(0.3))
    p = t.text_frame.paragraphs[0]
    r = p.add_run()
    _set_run(r, f"启研 · 第 3 次团队汇报   {n}/{total}", 11, False, MUTED)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.12), Inches(7.5))
    bar.line.fill.background()
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL


def slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    return s


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    s = slide(prs)
    add_bg(s)
    k = s.shapes.add_textbox(Inches(0.9), Inches(2.0), Inches(11), Inches(0.4))
    r = k.text_frame.paragraphs[0].add_run()
    _set_run(r, "智能化软件系统与工程 · 第 3 次团队汇报", 14, True, TEAL)
    t = s.shapes.add_textbox(Inches(0.9), Inches(2.5), Inches(11), Inches(1.1))
    r = t.text_frame.paragraphs[0].add_run()
    _set_run(r, "启研 · 本科生进科研", 40, True, INK)
    sub = s.shapes.add_textbox(Inches(0.9), Inches(3.7), Inches(11), Inches(1.4))
    tf = sub.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    _set_run(r, "刘弘雅　别克扎提·拜别提　陈浩文　张效端　陈旭", 18, False, INK)
    p = tf.add_paragraph()
    r = p.add_run()
    _set_run(r, "2026 年 10 月 12 日课上　不超过 5 分钟", 16, False, MUTED)
    footer(s, 1)

    s = slide(prs)
    kicker(s, "项目回顾")
    title(s, "本机链路没重做，这轮把四块补厚了")
    bullets(s, [
        "已经能走通：对话认识用户 → 方向树 → 出小作业 → 提交后记下来。",
        "任务 1 前端：粒子首页、侧栏五栏、文案收到 w38。线上就是这版。",
        "任务 2 交互：进门改成先聊几句，不再摊一张长问卷。旧五问还留在代码里。",
        "任务 3 项目：公开来源检索、交 zip、五条标准评阅。查不到就空着。",
        "任务 4 路径：认知 / 经济 / 数学 / AI 各一条入门路径，医学只写基础医学。",
        "没做成：书面问卷稿没有单独成册；方向树主干还没全部改回 FIELD_TREES。",
    ], size=17)
    footer(s, 2)

    s = slide(prs)
    kicker(s, "任务分工")
    title(s, "五个人，四块，对上任务表")
    bullets(s, [
        "任务 1 前端　刘弘雅　交：能打开的页面，工作区收干净。",
        "任务 2 交互问卷　别克扎提·拜别提　交：进门怎么聊，学完做完接下一件。",
        "任务 3 项目　陈浩文　交：来源、检索、压缩包要求、评分标准。",
        "任务 4 路径　张效端（认知 / 经济）、陈旭（数学 / AI）。",
        "任务 4 交：四条路径文本 + 改树建议，先不对同一份树抢着改代码。",
    ])
    footer(s, 3)

    s = slide(prs)
    kicker(s, "谁的结果给谁")
    title(s, "2 和 4 出草稿，1 和 3 再跟")
    bullets(s, [
        "任务 2 的交互 → 任务 1 按这个画界面，不先把流程画死。",
        "任务 4 的路径 → 任务 1 改树的样子；任务 3 按阶段去找项目。",
        "任务 3 的项目和评分 → 任务 2 接进「做完之后下一步问什么」。",
        "接口没另起一套：沿用现有 REST，要改先在组里说。",
    ])
    footer(s, 4)

    s = slide(prs)
    kicker(s, "系统部署")
    title(s, "开发 / 测试 / 生产分开")
    bullets(s, [
        "开发：本机。Dockerfile + compose，端口 8100。",
        "测试：本机 8101，或云上函数 qiyan-test。只给组里，不当首页。",
        "生产：GitHub Pages 页面 + 杭州函数 qiyan 做接口。",
        "选 Serverless：按量 ECS 开不了（后付费额度不够），不充值。",
        "容器：根目录 Dockerfile，控制台文件名就填 Dockerfile。",
        "数据：容器里在 /data；云函数写 /tmp，回收会丢。",
    ], size=17)
    footer(s, 5)

    s = slide(prs)
    kicker(s, "和讲义的差距")
    title(s, "讲义要服务器 + 域名，我们差在这两处")
    bullets(s, [
        "讲义：有公网 IP 的服务器，通常配域名。",
        "实际：函数计算，没有常驻机器。",
        "讲义：后端登记备案域名。",
        "实际：没有备案域名，不买、不备案。",
        "讲义：容器统一环境。实际：已补镜像；云上仍是函数自定义运行时。",
        "fcapp.run 当首页会下载 htm，所以用户入口用 GitHub Pages。",
    ], size=17)
    footer(s, 6)

    s = slide(prs)
    kicker(s, "系统发布")
    title(s, "打开就能用，不用装")
    bullets(s, [
        "渠道：GitHub Pages（页面）+ 阿里云函数计算（接口）。",
        "给用户的话：启研，给北大本科一、二年级。先聊你卡在哪，再给一件能做完的事。",
        "演示地址：https://chengcanwu.github.io/AutoResearch/",
        "打开后：点进入启研 → 起个称呼 → 先聊几句 → 看今日。",
        "不要点 fcapp.run 当首页，那个地址会下载文件。",
    ], size=17)
    footer(s, 7)

    s = slide(prs)
    kicker(s, "团队追踪")
    title(s, "先文本后代码，接口没擅自改")
    bullets(s, [
        "建议顺序发生了：任务 2、4 先出路径和交互说明，1、3 再往产品上收。",
        "任务 4 两人分开搜，合成同一份改树建议，再让前端动树。",
        "任务 3 的 paths.json 给项目检索用，和路径文本对齐。",
        "REST 没另起炉灶。密钥不进仓库。",
    ])
    footer(s, 8)

    s = slide(prs)
    kicker(s, "AI 辅助")
    title(s, "正面能加快，负面会把坑藏起来")
    bullets(s, [
        "正面：补 Dockerfile / 发布脚本；对照讲义写清三套环境；收 PPT 结构。",
        "正面：本机没有 Docker 时，用 GitHub Actions 留下 build / run 输出。",
        "负面：默认函数域名会强制下载网页，页面和接口只好拆开。",
        "负面：PowerShell 改中文文件会弄坏编码，前端只能用手写工具改。",
        "负面：AI 容易把「想做成」写成「已经做成」。没交的书面稿这页不编。",
    ], size=17)
    footer(s, 9)

    s = slide(prs)
    kicker(s, "演示")
    title(s, "现在就能打开")
    bullets(s, [
        "https://chengcanwu.github.io/AutoResearch/",
        "未完成：生产库不持久；没有备案域名；未开 ECS。",
        "书面问卷稿、方向树主干改代码，还要组里收一版。",
        "10 月 12 日前发布已经完成。",
    ])
    footer(s, 10)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print("wrote", OUT, "bytes", OUT.stat().st_size)


if __name__ == "__main__":
    main()
